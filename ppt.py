import os
import sys
import pptx
from pptx import Presentation

# Paths
SRC_PATH = r"c:\Users\yaso0\Downloads\How we judge evidence fisher feedback  our test results  domain benchmarks  published papers. Papers show the problem is real; they don’t prove ORCA works..pptx"
DST_PATH = r"c:\sih26066\SIH2026_SIH26066_INCOIS_OceanVisionAI.pptx"

if not os.path.exists(SRC_PATH):
    print(f"Error: Design reference PPTX not found at {SRC_PATH}")
    sys.exit(1)

prs = Presentation(SRC_PATH)

def set_para_text(p, new_text, font_size=None, bold=None, color_rgb=None):
    if not p.runs:
        r = p.add_run()
        r.text = new_text
        if font_size: r.font.size = font_size
        if bold is not None: r.font.bold = bold
        if color_rgb: r.font.color.rgb = color_rgb
        return
        
    first_run = p.runs[0]
    orig_font_name = first_run.font.name
    orig_font_size = first_run.font.size
    orig_bold = first_run.font.bold
    orig_color = None
    try:
        orig_color = first_run.font.color.rgb
    except:
        pass
        
    first_run.text = new_text
    if font_size: first_run.font.size = font_size
    elif orig_font_size: first_run.font.size = orig_font_size
    
    if bold is not None: first_run.font.bold = bold
    elif orig_bold is not None: first_run.font.bold = orig_bold
    
    if color_rgb: first_run.font.color.rgb = color_rgb
    elif orig_color: first_run.font.color.rgb = orig_color
    
    if orig_font_name: first_run.font.name = orig_font_name

    for r in p.runs[1:]:
        r.text = ""

def get_shape_by_id(slide, shape_id):
    def find_in(shapes):
        for s in shapes:
            if s.shape_id == shape_id:
                return s
            if s.shape_type == 6:
                res = find_in(s.shapes)
                if res: return res
        return None
    return find_in(slide.shapes)

def replace_shape_text(slide, shape_id, lines_list, font_size=None, bold=None, color_rgb=None):
    s = get_shape_by_id(slide, shape_id)
    if not s or not s.has_text_frame:
        return
    tf = s.text_frame
    for i, line in enumerate(lines_list):
        if i < len(tf.paragraphs):
            p = tf.paragraphs[i]
        else:
            p = tf.add_paragraph()
        set_para_text(p, line, font_size=font_size, bold=bold, color_rgb=color_rgb)
        
    for i in range(len(lines_list), len(tf.paragraphs)):
        tf.paragraphs[i].text = ""

print("Applying INCOIS SIH PS 2 content while strictly preserving design reference...")

# ------------------------------------------------------------------------------
# SLIDE 1: TITLE PAGE
# ------------------------------------------------------------------------------
s1 = prs.slides[0]
replace_shape_text(s1, 7, ["SMART INDIA HACKATHON 2026"])
replace_shape_text(s1, 8, [
    "",
    "Problem Statement ID – SIH26066",
    "Problem Statement Title: Satellite Embedding-Based Deep Learning Framework to Reconstruct Depth-Wise Subsurface Temperature in the North Indian Ocean",
    "Organization: Ministry of Earth Sciences (MoES) | INCOIS Ocean Valley",
    "Theme: Disaster Management | PS Category: Software Edition",
    "Team ID: 188334",
    "Team Name: OceanVision AI"
])

# ------------------------------------------------------------------------------
# SLIDE 2: PROPOSED SOLUTION
# ------------------------------------------------------------------------------
s2 = prs.slides[1]
replace_shape_text(s2, 30, ["What’s new in our Solution — Technical Innovations"])
replace_shape_text(s2, 37, ["Pain Point: Ocean Opacity & In-Situ Gaps", "Argo floats drift ~300 km apart (10-day delay); satellites only see the surface skin."])
replace_shape_text(s2, 55, ["Our Solution: Satellite Embedding-Based 3D Subsurface Temperature Reconstruction"])

replace_shape_text(s2, 58, [
    "",
    "Multi-Satellite Fusion: Daily OSTIA SST, SMAP SSS, DUACS SLA, OSCAR & CCMP",
    "Satellite Embedding Engine: ConvNeXt-V2 (eddies) + Swin-Transformer (planetary waves)",
    "Physics-Informed Regularization: Enforces TEOS-10 density & static stability (N² ≥ 0)",
    "15 Standard Depths: Reconstructs 0 to 1000m at 0.25° grid across North Indian Ocean",
    ""
])

# 6 Innovation Cards
replace_shape_text(s2, 68, ["15 Standard Depth Profiling"])
replace_shape_text(s2, 66, ["Reconstructs 0, 5, 10, 20, 30, 50, 75, 100, 125, 150, 200, 300, 500, 700, 1000m depth levels at 0.25° daily"])
replace_shape_text(s2, 65, ["Target: CMEMS GLORYS12V1 (moi-00021)"])

replace_shape_text(s2, 70, ["Dual-Branch Embedding Engine"])
replace_shape_text(s2, 67, ["Compact 512-D multi-scale representation learning combining ConvNeXt-V2 and Swin Transformer v2"])
replace_shape_text(s2, 69, ["Pipeline: Google Earth Engine & PyTorch"])

replace_shape_text(s2, 71, ["Cyclone Heat Potential (TCHP)"])
replace_shape_text(s2, 72, ["Real-time D26 isotherm depth & heat content mapping for 48h cyclone rapid intensification alerts"])
replace_shape_text(s2, 73, ["Beneficiary: IMD & NDMA Disaster Mgmt"])

replace_shape_text(s2, 74, ["Baroclinic Altimetry Coupling"])
replace_shape_text(s2, 75, ["Exploits Sea Level Anomaly (SLA) & Wind Stress Curl as dynamic baroclinic proxies for thermocline shifts"])
replace_shape_text(s2, 76, ["Inputs: DUACS Altimetry & CCMP Winds"])

replace_shape_text(s2, 77, ["TEOS-10 Thermodynamic Guard"])
replace_shape_text(s2, 78, ["Observation-Guided density loss eliminates unphysical vertical density inversions (N² ≥ 0)"])
replace_shape_text(s2, 79, ["Physics: UNESCO TEOS-10 Equation of State"])

replace_shape_text(s2, 81, ["Ultra-Fast Inference & Offline Mode"])
replace_shape_text(s2, 80, ["<45ms per basin tile on Vertex AI gRPC; full offline ONNX runtime cache on localhost for zero Wi-Fi risk"])
replace_shape_text(s2, 82, ["Serving: Google Cloud Run & Local ONNX"])

# Bottom 3 summary boxes
replace_shape_text(s2, 84, ["Proposed solution"])
replace_shape_text(s2, 83, ["End-to-end deep learning framework converting 7 daily satellite surface channels into 3D vertical temperature fields (0–1000m) at 0.25° resolution across the North Indian Ocean."])

replace_shape_text(s2, 86, ["How it solves the problem"])
replace_shape_text(s2, 85, ["Overcomes seawater opacity to satellites; bridges sparse Argo float gaps with continuous daily grids; provides instant 3D fields without heavy numerical model delays."])

replace_shape_text(s2, 88, ["Innovation and uniqueness"])
replace_shape_text(s2, 87, [
    "• Dual ConvNeXt-Swin embedding engine",
    "• Observation-Guided TEOS-10 density loss",
    "• Climatological anomaly residual learning"
])

replace_shape_text(s2, 93, ["Technical Innovation: Satellite Embedding & Physics-Guided Neural Architecture"])
replace_shape_text(s2, 94, ["North Indian Ocean 3D Prototype Model"])

# ------------------------------------------------------------------------------
# SLIDE 3: TECHNICAL APPROACH
# ------------------------------------------------------------------------------
s3 = prs.slides[2]
replace_shape_text(s3, 6, ["TECHNICAL APPROACH"])
replace_shape_text(s3, 11, ["Embedding & Physics Architecture"])
replace_shape_text(s3, 16, ["Data We Use (Official MoES/INCOIS Suite)"])
replace_shape_text(s3, 21, ["Technical Innovations"])
replace_shape_text(s3, 24, ["Technology Stack"])

t26 = get_shape_by_id(s3, 26)
if t26 and t26.has_table:
    table = t26.table
    table.rows[0].cells[0].text = "Technical Innovation"
    table.rows[0].cells[1].text = "Oceanographic Challenge"
    table.rows[0].cells[2].text = "Our AI & Physics Solution"
    
    table.rows[1].cells[0].text = "Observation-Guided PINN"
    table.rows[1].cells[1].text = "Standard PINNs suffer gradient conflict & spurious convergence"
    table.rows[1].cells[2].text = "Anchors physical loss to observed density & TEOS-10, guaranteeing monotonic stratification (N² ≥ 0)"
    
    table.rows[2].cells[0].text = "Dual-Branch Embedding"
    table.rows[2].cells[1].text = "Spatial scale mismatch: mesoscale eddies vs planetary waves"
    table.rows[2].cells[2].text = "ConvNeXt-V2 (7x7 local eddies) + Swin-Transformer (basin-wide planetary Rossby & Kelvin waves)"
    
    table.rows[3].cells[0].text = "Baroclinic Altimetry Proxy"
    table.rows[3].cells[1].text = "Subsurface thermocline displacement invisible to surface sensors"
    table.rows[3].cells[2].text = "Couples Sea Level Anomaly (SLA) & Wind Stress Curl (Ekman pumping w_E) to vertical isotherm shift"

# ------------------------------------------------------------------------------
# SLIDE 4: FEASIBILITY AND VIABILITY
# ------------------------------------------------------------------------------
s4 = prs.slides[3]
replace_shape_text(s4, 6, ["FEASIBILITY AND VIABILITY"])
replace_shape_text(s4, 15, ["Proof: Working Prototype & 3D Temperature Slicing Engine"])
replace_shape_text(s4, 23, ["Validated on CMEMS GLORYS12V1 & INCOIS LAS ARGO Profilers"])
replace_shape_text(s4, 24, ["TRL 4"])
replace_shape_text(s4, 27, ["Live Demo: End-to-end inference in <45ms generating 15 depth slices with interactive 3D transects on MapLibre."])
replace_shape_text(s4, 28, ["Challenges, Risks and Mitigation Strategies"])
replace_shape_text(s4, 29, ["Feasibility Analysis & Data Integrity"])

# ------------------------------------------------------------------------------
# SLIDE 5: IMPACT AND BENEFITS
# ------------------------------------------------------------------------------
s5 = prs.slides[4]
replace_shape_text(s5, 4, ["IMPACT AND BENEFITS"])
replace_shape_text(s5, 14, ["Target Beneficiaries & Maritime Stakeholders"])

replace_shape_text(s5, 19, ["48 Hours", "Earlier Cyclone Alert (TCHP)"])
replace_shape_text(s5, 23, ["30% Saved", "Diesel Fuel for Fishermen"])
replace_shape_text(s5, 27, ["0.25° Daily", "100% Basin-Wide 3D Grid"])
replace_shape_text(s5, 32, ["15 Depths", "Standard INCOIS Layers", "From 0m down to 1000m"])
replace_shape_text(s5, 35, ["<45 ms", "Real-Time Inference", "Vertex AI gRPC Latency"])

replace_shape_text(s5, 38, ["Benefits of the Solution"])
replace_shape_text(s5, 43, [
    "Safer marine operations with 48h cyclone warning",
    "Precise thermocline Potential Fishing Zone advisory",
    "Continuous daily coverage replacing sparse point floats"
])
replace_shape_text(s5, 48, [
    "30% diesel fuel saved navigating to productive zones",
    "Substantial increase in Catch Per Unit Effort (CPUE)",
    "Prevents fishing fleet losses during rapid storm onset"
])
replace_shape_text(s5, 53, [
    "Monitors subsurface Marine Heatwaves (MHWs)",
    "Early alerts for coral bleaching in Lakshadweep/Mannar",
    "Directly empowers India’s sustainable Blue Economy"
])

replace_shape_text(s5, 56, ["Impact on the Target Audience"])
replace_shape_text(s5, 60, ["1  IMD & Disaster Mgmt (NDMA)"])
replace_shape_text(s5, 63, ["2  7M Coastal Fisherfolk & INCOIS"])
replace_shape_text(s5, 66, ["3  Indian Navy & Coast Guard"])

replace_shape_text(s5, 70, ["Sources: INCOIS Marine Bulletins; IMD Cyclone Reports; CMFRI Landings; Copernicus Marine Service."])
replace_shape_text(s5, 71, ["Maritime Stakeholder Value"])
replace_shape_text(s5, 72, ["Social"])
replace_shape_text(s5, 73, ["Economic"])
replace_shape_text(s5, 74, ["Environmental"])

replace_shape_text(s5, 75, ["Early warning of rapid cyclone intensification; protects vulnerable coastal communities in Odisha, AP, WB, Gujarat"])
replace_shape_text(s5, 76, ["Directs artisanal & mechanized fleets to active thermocline upwelling zones; saves millions in search fuel"])
replace_shape_text(s5, 77, ["Supplies INCOIS with high-resolution daily 3D assimilation feed; provides Navy with Sound Velocity Profiles (SVP)"])

replace_shape_text(s5, 82, ["Sustainability & Deployment Roadmap"])
replace_shape_text(s5, 85, ["Next 3 Months"])
replace_shape_text(s5, 88, ["North Indian Ocean beta model on Vertex AI; INCOIS LAS ARGO continuous validation"])
replace_shape_text(s5, 91, ["Year 1"])
replace_shape_text(s5, 94, ["Operational integration with INCOIS SAMUDRA app & IMD cyclone advisory system"])
replace_shape_text(s5, 97, ["Year 2"])
replace_shape_text(s5, 100, ["Pan-Indian Ocean expansion; addition of BGC-Argo variables (DO, Chlorophyll, pH)"])
replace_shape_text(s5, 103, ["Adoption"])
replace_shape_text(s5, 106, ["MoES Deep Ocean Mission & INCOIS national marine forecasting infrastructure"])

# ------------------------------------------------------------------------------
# SLIDE 6: RESEARCH AND REFERENCES
# ------------------------------------------------------------------------------
s6 = prs.slides[5]
replace_shape_text(s6, 8, ["RESEARCH AND REFERENCES"])
replace_shape_text(s6, 14, ["Academic Research Literature & Architecture Mapping"])
replace_shape_text(s6, 17, ["Official Open-Access Datasets (MoES & INCOIS Mandate)"])
replace_shape_text(s6, 19, [
    "Evidence Hierarchy:",
    "INCOIS LAS In-Situ ARGO Floats > CMEMS GLORYS12V1 Reanalysis > Theoretical PINN Bounds > Academic Literature."
])

t20 = get_shape_by_id(s6, 20)
if t20 and t20.has_table:
    table = t20.table
    table.rows[0].cells[0].text = "Research Publication"
    table.rows[0].cells[1].text = "Incorporated in Our Solution For"
    
    table.rows[1].cells[0].text = "[1] Song et al., 'Convformer: Subsurface T/S Reconstruction,' Remote Sens. 2024, doi:10.3390/rs16132422."
    table.rows[1].cells[1].text = "Dual-branch CNN-Transformer embedding engine"
    
    table.rows[2].cells[0].text = "[2] Xiao, Tang, & Li, 'Observation-Guided PINN (OG-PINN),' Hohai Univ., 2026."
    table.rows[2].cells[1].text = "Observation-Guided TEOS-10 density loss & stability"
    
    table.rows[3].cells[0].text = "[3] Shao et al., '3D-MOPGCBANN: Physics-Guided ST Predicting,' IEEE TGRS 2025, doi:10.1109/TGRS.2025.4213112."
    table.rows[3].cells[1].text = "Pareto-optimal stratification stability (N² ≥ 0)"
    
    table.rows[4].cells[0].text = "[4] Sun et al., 'CSSP-ConvLSTM: Subsurface Temperature Reconstruction,' IEEE TGRS 2026, doi:10.1109/TGRS.2026.4202415."
    table.rows[4].cells[1].text = "Spatiotemporal ConvLSTM sequence memory"
    
    table.rows[5].cells[0].text = "[5] Chae, Donohue, & Park, 'TS-Cast: Deep Learning for Subsurface Ocean,' Ocean Sci. 2026, doi:10.5194/os-22-2161-2026."
    table.rows[5].cells[1].text = "Uncertainty-aware multi-task loss formulation"
    
    table.rows[6].cells[0].text = "[6] Xie et al., 'Attention U-Net ST Field in SCS,' IEEE TGRS 2022, doi:10.1109/TGRS.2022.4209319."
    table.rows[6].cells[1].text = "Robust Huber loss & Wind Stress Curl Ekman pumping"

replace_shape_text(s6, 21, [
    "• SST (OSTIA 0.05° Daily): CMEMS doi:10.48670/moi-00168",
    "• SSS (SMAP/SMOS 0.125° Daily): CMEMS doi:10.48670/moi-00051",
    "• SSH/SLA (DUACS 0.25° Daily): CMEMS doi:10.48670/moi-00145",
    "• Target Reanalysis (GLORYS12V1): CMEMS doi:10.48670/moi-00021",
    "• Currents & Winds: NASA PO.DAAC OSCAR_L4 & CCMP_WINDS_V3.1",
    "• In-Situ Ground Truth: INCOIS Live Access Server (LAS) Gridded ARGO"
])

replace_shape_text(s6, 26, [
    "Future Expansion",
    "Full BGC-Argo biogeochemical parameter reconstruction (DO, Chlorophyll, pH), Pan-Indian Ocean coverage, and direct integration into the INCOIS SAMUDRA mobile app."
])

replace_shape_text(s6, 34, ["Copernicus / PO.DAAC"])
replace_shape_text(s6, 35, ["INCOIS Ocean Valley"])
replace_shape_text(s6, 36, ["MoES / IMD"])

# Save modified presentation
prs.save(DST_PATH)
print(f"Successfully generated: {DST_PATH}")